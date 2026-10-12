"""Step 1: Compile, geocode and footprint Link REIT's Hong Kong properties.

Full run: all 130 HK assets (`make assets`). `--pilot`: the 5 AXA assets in
config/assumptions.yaml `pilot`. Assets in footprint.point_only_valuation_nos (no. 130,
development site) get their ALS point only: no footprint, levels or flags.

Inputs:
    data/raw/link_valuation_2026-03-31.xlsx, sheet "Summary Report 260331": valuation
        no. and property name (columns found by header text; HK section ends at
        "Mainland China Properties"), plus the "Value as at" column only to order the
        review list. Values for losses, areas and income are step 2.
    data/manual/asset_name_aliases.csv: valuation no. -> iB1000 building names (rows whose
        alias_source starts with "iB1000"), with evidence; also the lookup aliases tried
        for an automatic exact seed match (see Footprints).
    data/interim/geocode_accepted.csv: accepted ALS point(s) per asset, written by
        scripts/geocode_als.py --parse from the raw ALS cache (data/raw/geocode_cache/),
        data/manual/als_points.csv and data/manual/geocode_accepted_overrides.csv.
        Priority: manual > ALS query passing the name check > reviewer override
        (decision 2026-10-06). Any asset without an accepted point stops the run. This
        step never calls the network. Its `district` column (ALS district, or for manual
        points the district in als_points.csv, else blank) is the asset's district.
    data/manual/footprint_overrides.csv: footprint decisions. action 'include' | 'exclude'
        per BUILDINGID; 'seed' per BUILDINGID makes that block a seed (hand footprint for an
        asset without a named iB1000 seed, decision 2026-10-07; seed_source 'override_seed');
        'accept' (blank buildingid) records sign-off of an asset's
        footprint without changing it; 'accept_location' (blank buildingid) the same for an
        asset whose location itself is approximate (review_status 'location_approximate',
        decision 2026-10-07). decided_by 'rule_v1' rows come from
        scripts/footprint_rules.py; they do not count as hand QA (qa_reviewed).
    data/raw/ib1000/ib1000_tile_index.gpkg: layer TileIndex (SHEETNO, REVISIONDATE).
    data/raw/ib1000/sheets/<sheet>/<sheet>.gdb: iB1000 FGDB per sheet.
    data/raw/lidar_dtm_2020/<sheet without dashes>(...)/*.tif: CEDD 2020 LiDAR DTM, 0.5 m,
        mPD (same 750 x 600 m grid as iB1000).
    Sheets and tiles are downloaded by scripts/fetch_terrain.py (never by this step).

Outputs (EPSG:2326; EPSG:4326 only for the site export in step 7):
    data/interim/ib1000_extract.gpkg: layers Building, Site, SpotHeight, BUILDINGNAME,
        BUILDINGRELATEBUILDINGNAME, SITENAME, SITERELATESITENAME; every row carries
        sheet_no and sheet_revision_date (tile index REVISIONDATE). Rebuilt each run.
    data/interim/assets.geojson: one row per ALS point; als_source = manual | als | accepted.
        shared_point (also in asset_footprints.geojson): the asset's accepted point is
        identical to another asset's (shared_point_nos(); nos. 104 and 107).
    data/interim/asset_footprints.geojson: one row per asset (footprint, flags, levels);
        footprint_status = matched | no_named_seed | point_only (no geometry unless matched);
        footprint_method = ib1000_footprint | point_buffer_approx (levels from a circle of
        elevation.approx_point_buffer_m around the accepted point(s): entry_level_mpd = p10,
        ring_median_mpd = median of the circle; no slope/podium flags).
    data/interim/ib1000_sheets_needed.csv: sheets for the extended seed search that are
        not downloaded (sheet_no, for_valuation_nos); the run stops if any.
    data/interim/lidar_tiles_needed.csv: LiDAR tiles under each footprint + ring_width_m,
        or under the approx circle (sheet_no, tile, have). Written before levels; the run
        stops if a tile is missing.
    data/interim/seed_suggestions.csv: for each no_named_seed asset, iB1000 names that
        contain the asset's estate token AND one of SHOPPING CENTRE, COMMERCIAL CENTRE,
        CAR PARK, MARKET, PLAZA, MALL ("Ctr", "Carpark" expanded), on blocks within
        seed_search_radius_m of an accepted point. Token: the derived lookup alias
        ('derived_alias') or the valuation name with the same prefix/suffix rules
        ('valuation_name_stripped'); "(I)".."(IV)" and trailing Estate/Court dropped.
        A regenerated view: its decision column is read from data/manual/seed_decisions.csv
        (tracked; valuation_no, candidate_name, buildingids, decision yes|no, decided_by,
        date). Decisions are recorded there, never in the view (step 1 stops if the view has
        a decision the tracked file lacks). scripts/apply_seed_suggestions.py turns 'yes'
        rows into alias rows (make apply-seeds).
    data/interim/asset_blocks.geojson: one row per footprint block with its reason and
        the QA suggestion for touching unnamed blocks.
    data/interim/touching_review.csv: for the review list, one row per undecided touching
        unnamed block (valuation_no, name, buildingid, type, base_mpd, roof_mpd, area_m2,
        suggestion, reason = one line from touching_reason(), decision blank), so the user
        can decide by exception. Decisions typed there are kept when the file is rewritten,
        but only data/manual/footprint_overrides.csv changes a footprint: `make
        apply-touching` copies them there (scripts/apply_touching_review.py;
        ACCEPT_SUGGESTIONS=1 also applies the remaining 'include' suggestions).
    data/interim/asset_review_status.csv: every asset with priority and review_status
        (point_only | location_approximate | footprint_unconfirmed | review | approx_pending_week3 | auto_accepted_hillside |
        hand_reviewed | accepted_clean; see review_status()); also columns of
        asset_footprints.geojson.
    data/interim/asset_triage.csv: the review list (decisions 2026-10-06): review_status
        'review' = priority A with at least one review reason; informational reasons (INFO_REASONS:
        elevation:slope, elevation:district_outlier, geocode:low_score) go to info_flags
        and alone do not put an asset on the list. priority A = entry_level_mpd (approx
        for unmatched assets) < triage.priority_a_below_mpd or unknown, else B.
        needs_decision = no_named_seed or touching_unnamed_undecided. Sorted by entry
        level ascending, then valuation value descending (value_hkd, internal only).
        Reasons (config `triage`):
        geocode:low_score (info) / geocode:phase_mismatch  accepted ALS point flagged
        geocode:als_far_from_footprint   an ALS point > als_max_distance_m from the footprint
        footprint:no_named_seed          no seed block found; see seed_suggestions.csv
        footprint:touching_unnamed_undecided  touching unnamed block(s) with no include /
            exclude row in footprint_overrides.csv and no 'accept' row for the asset
        footprint:edge_touch             a block touches the outer edge of the loaded sheets
        elevation:slope (info) / elevation:podium   slope_flag / podium_flag
        elevation:ring_outside_lidar     the ring is not fully covered by the LiDAR tiles
        elevation:no_level               no ground level from LiDAR (or the 5 m fallback):
            stays on the review list (needs a decision)
        elevation:district_outlier (info)  entry_level_mpd outside [Q1 - k*IQR, Q3 + k*IQR] of
            our assets in the same ALS district (districts with >= min_assets levels)

Footprints (decision 2026-10-01; replaces the 2026-09-24 Site-based spec):
    The iB1000 Site layer holds estate boundaries (SITECODE 'EST', 18-30 blocks incl.
    residential towers), not lot boundaries, so it is NOT used for the footprint. It is
    kept only as a context column (estate_sites).
    1. Sheets: every iB1000 sheet within sheet_search_buffer_m of the asset's ALS points.
       If no seed is found there, the search is repeated over the sheets within
       seed_search_radius_m (seed_search_m records which). Pieces clipped at sheet edges
       are unioned by BUILDINGID once, over all loaded sheets.
    2. Seeds: Building blocks in the asset's sheets linked (BUILDINGRELATEBUILDINGNAME) to
       a seed name. Seed names (seed_source): the asset's iB1000 alias rows if any
       ('alias_ib1000'); otherwise iB1000 English names equal to the valuation name or one
       of its ALS-sourced aliases after normalising case, spacing and punctuation
       ('auto_exact', decision 2026-10-06). Derived lookup aliases are never seeds. No
       fuzzy matching: anything else is footprint:no_named_seed.
    3. Add UNNAMED blocks touching a seed (within touch_tolerance_m), e.g. a podium.
       Named blocks of other buildings are never added.
    4. Apply data/manual/footprint_overrides.csv (include / exclude by BUILDINGID; user
       rows and rule_v1 rows from scripts/footprint_rules.py alike).
    5. QA suggestion for each touching unnamed block: 'exclude' if its BASELEVEL is more
       than suggest_exclude_above_seed_m above the seeds' min BASELEVEL (sits on the
       podium), 'review' if it has no BASELEVEL, else 'include'. A suggestion only; the
       footprint changes only through the overrides file.
    6. edge_touch: a block touching the outer edge of all loaded sheets may continue on a
       sheet not loaded.
    Automatic matching aims at >= 90% of assets (quality target); failures are
    hand-drawn during QA, not dropped.

Footprint area is a QA column only. Replacement cost (step 2) uses reported GFA or
IFA / efficiency_ratio, never footprint area.

Elevation (decisions 2026-09-24, 2026-10-06; config `elevation`):
    The ring is built from the FINAL footprint, i.e. after footprint_overrides.csv
    (footprint_levels). ring = footprint buffered by ring_width_m minus the footprint;
    LiDAR pixels selected by pixel centre; percentiles use linear interpolation.
    entry_level_mpd = entry_level_percentile of the ring: whether water enters.
        ring_median_mpd is the sensitivity; ring_min_mpd and ring_p90_mpd are QA.
    lowest_floor_mpd = min iB1000 BASELEVEL of the final footprint blocks: depth inside
        once water enters (used in Week 3). Same value as min_baselevel_mpd, which
        drives podium_flag.
    podium_flag = min BASELEVEL - ring median > podium_flag_threshold_m.
    slope_flag = ring p90 - ring p10 > slope_flag_threshold_m (sloping site, one ring
        level may not fit every entrance). Temple Mall North (slope choice) is deferred to
        Week 3 (rain score); for coastal flooding it is irrelevant (far above any modelled sea level).
    iB1000 BASELEVEL is not used for the ground level outside the building.

iB1000 at full scale (decision 2026-09-24): raw sheet zips may be deleted after
extraction into ib1000_extract.gpkg (data/SOURCES.md keeps URL + sha256).
"""

import argparse
import json
import re
import sys
import warnings
from pathlib import Path
from urllib.parse import quote

import geopandas as gpd
import numpy as np
import openpyxl
import pandas as pd
import pyogrio
import rasterio
import rasterio.windows
from rasterio.features import geometry_mask
from rasterio.merge import merge
from shapely.geometry import Point, box
from shapely.ops import unary_union

if __package__ in (None, ""):  # run as a script: make `pipeline` importable
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pipeline.config import load_config  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
MANUAL = ROOT / "data" / "manual"
INTERIM = ROOT / "data" / "interim"
XLSX = RAW / "link_valuation_2026-03-31.xlsx"
VALUATION_SHEET = "Summary Report 260331"
TILE_INDEX = RAW / "ib1000" / "ib1000_tile_index.gpkg"
SHEETS_DIR = RAW / "ib1000" / "sheets"
LIDAR_DIR = RAW / "lidar_dtm_2020"
LIDAR_INDEX = LIDAR_DIR / "dtm2020_grid_index.geojson"
DTM_5M = RAW / "dtm" / "dtm_5m.tif"
SEED_DECISIONS = MANUAL / "seed_decisions.csv"
SEED_DECISION_COLS = ["valuation_no", "candidate_name", "buildingids", "decision", "decided_by", "date"]

EXTRACT_LAYERS = {
    "Building": True, "Site": True, "SpotHeight": True,
    "BUILDINGNAME": False, "BUILDINGRELATEBUILDINGNAME": False,
    "SITENAME": False, "SITERELATESITENAME": False,
}
BLOCK_ATTRS = ["TYPEOFBUILDINGBLOCK", "BASELEVEL", "ROOFLEVEL", "BASELEVELDATASOURCE",
               "ROOFLEVELDATASOURCE", "BUILDINGSTATUS", "CERTAINTY", "LASTUPDATEDATE"]

warnings.filterwarnings("ignore", message="Measured")


# --- asset list and ALS points ---------------------------------------------------------

def read_hk_asset_names(xlsx: Path, sheet: str = VALUATION_SHEET) -> pd.DataFrame:
    """Return valuation no. and property name for every HK row.

    Each sub-table has its own header row; columns are found by the header text
    "No." / "Property Name". Reading stops at "Mainland China Properties".
    """
    ws = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)[sheet]
    rows, cols = [], None
    for row in ws.iter_rows(values_only=True):
        texts = [v.strip() if isinstance(v, str) else v for v in row]
        if "Mainland China Properties" in texts:
            break
        if "No." in texts and "Property Name" in texts:
            cols = (texts.index("No."), texts.index("Property Name"))
            continue
        if cols and isinstance(row[cols[0]], int) and isinstance(row[cols[1]], str):
            rows.append({"no": row[cols[0]], "name": row[cols[1]].strip()})
    return pd.DataFrame(rows)


def read_hk_asset_values(xlsx: Path, sheet: str = VALUATION_SHEET) -> pd.Series:
    """valuation no. -> value as at 31 March 2026 (HKD; the cells are plain HKD despite the
    "(HKD Million)" header, decision 2026-09-24). Column found by the header text starting
    "Value as at" in each sub-table. Used only to order the review list (step 1); step 2
    owns values."""
    ws = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)[sheet]
    out, cols = {}, None
    for row in ws.iter_rows(values_only=True):
        texts = [" ".join(v.split()) if isinstance(v, str) else v for v in row]
        if "Mainland China Properties" in texts:
            break
        if "No." in texts and "Property Name" in texts:
            val = [i for i, t in enumerate(texts) if isinstance(t, str) and t.startswith("Value as at")]
            if len(val) != 1:
                raise ValueError(f"expected one 'Value as at' column, header {texts}")
            cols = (texts.index("No."), val[0])
            continue
        if cols and isinstance(row[cols[0]], int) and isinstance(row[cols[1]], (int, float)):
            out[row[cols[0]]] = float(row[cols[1]])
    return pd.Series(out, name="value_hkd")


def als_cache_path(query: str, cache_dir: Path) -> Path:
    return cache_dir / f"{quote(query, safe='')}.json"


def als_lookup(query: str, cache_dir: Path) -> dict:
    """Return the cached raw ALS JSON response for `query`. Never calls the network."""
    path = als_cache_path(query, cache_dir)
    if not path.exists():
        raise FileNotFoundError(
            f"no cached ALS response for {query!r} at {path}. ALS is not queried from this "
            "pipeline; save the raw JSON response there (see data/SOURCES.md).")
    return json.loads(path.read_text())


def _as_list(x):
    return x if isinstance(x, list) else [x]


def _leaves(x) -> list[str]:
    if isinstance(x, dict):
        return [s for v in x.values() for s in _leaves(v)]
    if isinstance(x, list):
        return [s for v in x for s in _leaves(v)]
    return [str(x)] if x not in (None, "") else []


def parse_als_response(response: dict) -> pd.DataFrame:
    """Candidates tied at the top ALS score: easting, northing, lat, lon, score, address.

    Written against the ALS JSON layout (SuggestedAddress / Address / PremisesAddress /
    GeospatialInformation, ValidationInformation / Score); not yet checked against a real
    cached response, since ALS was unreachable on 2026-10-01.
    """
    rows = []
    for s in _as_list(response.get("SuggestedAddress", [])):
        prem = s["Address"]["PremisesAddress"]
        geo = _as_list(prem["GeospatialInformation"])[0]
        rows.append({
            "easting": float(geo["Easting"]), "northing": float(geo["Northing"]),
            "lat": float(geo["Latitude"]), "lon": float(geo["Longitude"]),
            "score": float(s["ValidationInformation"]["Score"]),
            "matched_address": ", ".join(_leaves(prem.get("EngPremisesAddress", {}))),
        })
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    return df[df.score == df.score.max()].reset_index(drop=True)


TRIAGE_GEOCODE_FLAGS = ("low_score", "phase_mismatch")


def geocode_triage(points: pd.DataFrame) -> pd.DataFrame:
    """One row per asset whose accepted point(s) carry a TRIAGE_GEOCODE_FLAGS flag:
    valuation_no, name, reasons ('geocode:<flag>', ';'-joined), als_source."""
    rows = []
    for no, pts in points.groupby("valuation_no", sort=True):
        flags = {f for s in pts["flags"].fillna("") for f in str(s).split(";") if f}
        hit = [f"geocode:{f}" for f in TRIAGE_GEOCODE_FLAGS if f in flags]
        if hit:
            rows.append({"valuation_no": no, "name": pts.name.iloc[0], "reasons": ";".join(hit),
                         "als_source": ",".join(sorted(set(pts.als_source)))})
    return pd.DataFrame(rows, columns=["valuation_no", "name", "reasons", "als_source"])


def load_accepted_points(assets: pd.DataFrame, accepted_csv: Path,
                         crs: str) -> gpd.GeoDataFrame:
    """Accepted ALS point(s) per asset from data/interim/geocode_accepted.csv (written by
    scripts/geocode_als.py --parse: manual > name-checked ALS > reviewer override).
    Stops, listing them, if any asset has no accepted point."""
    if not accepted_csv.exists():
        raise FileNotFoundError(f"{accepted_csv} missing: run scripts/geocode_als.py --parse")
    acc = pd.read_csv(accepted_csv, dtype={"query_used": str})
    lacking = assets[~assets.no.isin(acc.valuation_no)]
    if len(lacking):
        lst = "; ".join(f"no. {a.no} {a.name!r}" for a in lacking.itertuples())
        raise ValueError(f"{len(lacking)} asset(s) without an accepted geocode "
                         f"(see scripts/geocode_als.py --parse): {lst}")
    df = acc[acc.valuation_no.isin(assets.no)].drop(columns="name").merge(
        assets.rename(columns={"no": "valuation_no"})[["valuation_no", "name"]], on="valuation_no")
    df = df.rename(columns={"query_used": "query", "source": "als_source"})
    if "district" not in df:
        raise ValueError(f"{accepted_csv} has no district column: rerun scripts/geocode_als.py --parse")
    cols = ["valuation_no", "name", "query", "matched_address", "district", "easting", "northing",
            "lat", "lon", "score", "als_source", "flags", "provenance"]
    df = df[cols].sort_values(["valuation_no", "easting", "northing"]).reset_index(drop=True)
    return gpd.GeoDataFrame(df, geometry=gpd.points_from_xy(df.easting, df.northing), crs=crs)


def shared_point_nos(points: pd.DataFrame) -> set[int]:
    """Valuation nos. whose accepted point has exactly the same easting/northing as another
    asset's (nos. 104 and 107 matched one ALS estate address; decision 2026-10-06). Their
    point-based inputs are not independent; separate coordinates pending."""
    n = points.groupby(["easting", "northing"]).valuation_no.transform("nunique")
    return set(points.loc[n > 1, "valuation_no"])


# --- iB1000 extract ----------------------------------------------------------------------

def sheets_for_points(points, tile_index: gpd.GeoDataFrame, buffer_m: float,
                      sheets_dir: Path | None = None) -> pd.DataFrame:
    """Sheets (SHEETNO, REVISIONDATE) within buffer_m of any point. Raises if one is missing."""
    area = unary_union([Point(p).buffer(buffer_m) for p in points])
    hit = tile_index[tile_index.intersects(area)]
    out = pd.DataFrame({"sheet_no": hit.SHEETNO.values,
                        "sheet_revision_date": hit.REVISIONDATE.astype(str).values})
    if sheets_dir is not None:
        missing = [s for s in out.sheet_no if not (sheets_dir / s / f"{s}.gdb").exists()]
        if missing:
            raise FileNotFoundError(f"iB1000 sheets not downloaded: {missing}")
    return out.sort_values("sheet_no").reset_index(drop=True)


def extract_ib1000(sheets: pd.DataFrame, sheets_dir: Path, out_path: Path, crs: str) -> None:
    """Copy the listed layers of each sheet into one GeoPackage, tagged with sheet number
    and revision date. Overwrites out_path."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.unlink(missing_ok=True)
    for layer, spatial in EXTRACT_LAYERS.items():
        parts = []
        for s in sheets.itertuples():
            df = pyogrio.read_dataframe(sheets_dir / s.sheet_no / f"{s.sheet_no}.gdb",
                                        layer=layer, read_geometry=spatial)
            parts.append(df.assign(sheet_no=s.sheet_no, sheet_revision_date=s.sheet_revision_date))
        df = pd.concat(parts, ignore_index=True)
        if spatial:
            df = gpd.GeoDataFrame(df, geometry="geometry", crs=crs)
            df["geometry"] = df.geometry.force_2d()
        pyogrio.write_dataframe(df, out_path, layer=layer, driver="GPKG", append=out_path.exists())


def read_extract(path: Path, layer: str) -> pd.DataFrame:
    spatial = EXTRACT_LAYERS[layer]
    return pyogrio.read_dataframe(path, layer=layer, read_geometry=spatial)


# --- footprint ---------------------------------------------------------------------------

def merge_sheet_pieces(gdf: gpd.GeoDataFrame, id_col: str, attrs: list[str]) -> gpd.GeoDataFrame:
    """Union pieces of one feature clipped at sheet edges. Attributes must agree across
    pieces; adds `pieces` (number of sheets) and `sheets` (comma-joined). Features with a
    single piece pass through without a union (full run: ~10^5 blocks)."""
    multi = gdf[id_col].duplicated(keep=False)
    many = gdf[multi]
    for col in attrs:
        n = many.groupby(id_col)[col].nunique(dropna=False)
        if (n > 1).any():
            raise ValueError(f"{col} differs across sheet pieces for {id_col} {list(n[n > 1].index)}")
    cols = [id_col, "geometry", *attrs, "sheets"]
    single = gdf[~multi].rename(columns={"sheet_no": "sheets"})[cols]
    parts = [single]
    if len(many):
        agg = {c: "first" for c in attrs}
        agg["sheet_no"] = lambda x: ",".join(sorted(set(x)))
        parts.append(many.dissolve(by=id_col, aggfunc=agg).reset_index()
                     .rename(columns={"sheet_no": "sheets"})[cols])
    merged = gpd.GeoDataFrame(pd.concat(parts, ignore_index=True), geometry="geometry", crs=gdf.crs)
    merged["pieces"] = merged.sheets.str.count(",") + 1
    return merged.sort_values(id_col).reset_index(drop=True)


def building_names(names: pd.DataFrame, relate: pd.DataFrame) -> pd.Series:
    """BUILDINGID -> sorted tuple of English names (all sheets)."""
    n = names.drop_duplicates("BUILDINGNAMEID")[["BUILDINGNAMEID", "ENGLISHBUILDINGNAME"]]
    r = relate.drop_duplicates(["BUILDINGNAMEID", "BUILDINGID"])[["BUILDINGNAMEID", "BUILDINGID"]]
    j = r.merge(n, on="BUILDINGNAMEID").dropna(subset=["ENGLISHBUILDINGNAME"])
    return j.groupby("BUILDINGID").ENGLISHBUILDINGNAME.agg(lambda x: tuple(sorted(set(x))))


def norm_name(s: str) -> str:
    """Upper case, every run of non-alphanumerics (incl. '&', '*', brackets) -> one space."""
    return re.sub(r"[^0-9A-Z]+", " ", str(s).upper()).strip()


def seed_candidates(no: int, valuation_name: str, aliases: pd.DataFrame) -> tuple[list[str], str]:
    """Names that make an iB1000 block a seed, and their source.

    1. iB1000 alias rows (alias_source starts with "iB1000"): exact iB1000 names, used as
       given -> 'alias_ib1000'.
    2. Otherwise the valuation name and the asset's ALS-sourced aliases, compared with
       iB1000 English names after norm_name -> 'auto_exact'. No fuzzy matching. Aliases
       "derived from valuation name" are ALS queries only (decision 2026-10-06): mostly
       estate names, which iB1000 also gives to residential blocks.
    """
    a = aliases[aliases.valuation_no == no]
    src = a.alias_source.fillna("")
    ib = a[src.str.startswith("iB1000")]
    if len(ib):
        return sorted(set(ib.alias)), "alias_ib1000"
    als = a[src.str.startswith("HK Address Lookup Service")]
    return sorted({valuation_name, *als.alias}), "auto_exact"


def seed_names_for(candidates: list[str], how: str, names_by_id: pd.Series) -> set[str]:
    """iB1000 English names (as spelled in iB1000) that count as seeds for these candidates."""
    have = {n for tup in names_by_id.values for n in tup}
    if how == "alias_ib1000":  # as spelled, ignoring leading/trailing spaces in iB1000
        given = {c.strip() for c in candidates}
        return set(candidates) | {n for n in have if n.strip() in given}
    want = {norm_name(c) for c in candidates} - {""}
    return {n for n in have if norm_name(n) in want}


NAME_PREFIXES = ("Retail and Car Park within", "Car Park within", "Retail within")
NAME_SUFFIXES = ("Commercial Centre", "Shopping Centre", "Commercial Complex", "Plaza")
SUGGEST_KEYWORDS = ("SHOPPING CENTRE", "COMMERCIAL CENTRE", "CAR PARK", "MARKET", "PLAZA", "MALL")
ABBREVIATIONS = {"CTR": "CENTRE", "CARPARK": "CAR PARK"}  # iB1000 spellings, whole words


def strip_name(name: str) -> str:
    """Valuation name with a known prefix ("Car Park within" ...) or suffix ("Shopping
    Centre" ...) and a trailing '*' removed; the rule behind the derived lookup aliases."""
    t = " ".join(str(name).split()).rstrip("*").strip()
    for p in NAME_PREFIXES:
        if t.lower().startswith(p.lower() + " "):
            t = t[len(p):].strip()
            break
    for x in NAME_SUFFIXES:
        if t.lower().endswith(" " + x.lower()):
            t = t[: -len(x)].strip()
            break
    return t


def estate_token(text: str) -> str:
    """Estate-name token: phase markers "(I)".."(IV)" dropped, normalised (norm_name), then
    trailing ESTATE / COURT words dropped. 'Kin Ming Estate' -> 'KIN MING'."""
    t = norm_name(re.sub(r"\((I|II|III|IV)\)", " ", str(text), flags=re.I))
    words = t.split()
    while words and words[-1] in ("ESTATE", "COURT"):
        words.pop()
    return " ".join(words)


def asset_token(no: int, valuation_name: str, aliases: pd.DataFrame) -> tuple[str, str]:
    """(token, token_source): from the asset's derived lookup alias if it has one
    ('derived_alias'), else from the valuation name stripped the same way
    ('valuation_name_stripped')."""
    a = aliases[(aliases.valuation_no == no)
                & aliases.alias_source.fillna("").str.startswith("derived from valuation name")]
    if len(a):
        return estate_token(a.alias.iloc[0]), "derived_alias"
    return estate_token(strip_name(valuation_name)), "valuation_name_stripped"


def expand_abbrev(normed: str) -> str:
    return " ".join(ABBREVIATIONS.get(w, w) for w in normed.split())


def suggestion_keyword(name: str) -> str | None:
    """First SUGGEST_KEYWORDS phrase in the normalised, abbreviation-expanded name."""
    t = f" {expand_abbrev(norm_name(name))} "
    return next((k for k in SUGGEST_KEYWORDS if f" {k} " in t), None)


def seed_suggestions(blocks: gpd.GeoDataFrame, names_by_id: pd.Series, pts, token: str,
                     within_m: float) -> pd.DataFrame:
    """iB1000 names containing the token AND a SUGGEST_KEYWORDS phrase, on blocks within
    within_m of any point. One row per name: candidate_name, buildingids (all blocks with
    that name in range), sheets, distance_m (nearest block), keyword. Hints only."""
    cols = ["candidate_name", "buildingids", "sheets", "distance_m", "keyword"]
    tok = f" {token} " if token else None
    if not tok:
        return pd.DataFrame(columns=cols)
    area = unary_union([p.buffer(within_m) for p in pts])
    near = blocks.iloc[blocks.sindex.query(area, predicate="intersects")]
    near = near[near.BUILDINGID.isin(names_by_id.index)]
    hits = {}
    for r in near.itertuples():
        d = min(r.geometry.distance(p) for p in pts)
        if d > within_m:
            continue
        for n in names_by_id[r.BUILDINGID]:
            kw = suggestion_keyword(n)
            if kw and tok in f" {expand_abbrev(norm_name(n))} ":
                h = hits.setdefault(n.strip(), {"ids": set(), "sheets": set(), "d": np.inf, "kw": kw})
                h["ids"].add(int(r.BUILDINGID))
                h["sheets"].update(str(r.sheets).split(","))
                h["d"] = min(h["d"], d)
    rows = [{"candidate_name": n, "buildingids": ",".join(map(str, sorted(h["ids"]))),
             "sheets": ",".join(sorted(h["sheets"])), "distance_m": round(float(h["d"]), 1),
             "keyword": h["kw"]} for n, h in hits.items()]
    return pd.DataFrame(rows, columns=cols).sort_values(["distance_m", "candidate_name"]).reset_index(drop=True)


def read_seed_decisions(path: Path = SEED_DECISIONS) -> pd.DataFrame:
    """data/manual/seed_decisions.csv (tracked): the user's yes/no per suggestion."""
    if not path.exists():
        return pd.DataFrame(columns=SEED_DECISION_COLS)
    d = pd.read_csv(path, dtype=str, keep_default_na=False)
    d["valuation_no"] = d.valuation_no.astype(int)
    bad = set(d.decision.str.strip().str.lower()) - {"yes", "no"}
    if bad:
        raise ValueError(f"{path.name}: decision must be yes or no, found {sorted(bad)}")
    return d


def unrecorded_view_decisions(view: pd.DataFrame | None, decisions: pd.DataFrame) -> pd.DataFrame:
    """Rows of an existing seed_suggestions.csv whose non-blank decision is not the one in
    seed_decisions.csv (e.g. typed into the regenerated view). Step 1 stops on these rather
    than overwrite them."""
    if view is None or view.empty or "decision" not in view:
        return pd.DataFrame(columns=["valuation_no", "candidate_name", "decision"])
    v = view[view.decision.fillna("").str.strip() != ""][["valuation_no", "candidate_name", "decision"]]
    rec = dict(zip(zip(decisions.valuation_no, decisions.candidate_name),
                   decisions.decision.str.strip().str.lower()))
    keep = [rec.get((n, c)) != d.strip().lower() for n, c, d in v.itertuples(index=False)]
    return v[keep]


def merge_decisions(new: pd.DataFrame, old: pd.DataFrame | None) -> pd.DataFrame:
    """Carry decisions (from data/manual/seed_decisions.csv) onto the new suggestion rows,
    matched on valuation_no + candidate_name. Rows without a decision get a blank one."""
    new = new.drop(columns="decision", errors="ignore")
    if old is None or old.empty or "decision" not in old:
        return new.assign(decision="")
    keep = old[["valuation_no", "candidate_name", "decision"]].drop_duplicates(
        ["valuation_no", "candidate_name"], keep="last")
    out = new.merge(keep, on=["valuation_no", "candidate_name"], how="left")
    out["decision"] = out.decision.fillna("")
    return out


def open_suggestions(sugg: pd.DataFrame) -> pd.Series:
    """valuation_no -> number of suggestions not marked decision 'no' (case-insensitive)."""
    keep = sugg[sugg.decision.fillna("").str.strip().str.lower() != "no"]
    return keep.groupby("valuation_no").size()


def select_footprint_blocks(blocks: gpd.GeoDataFrame, names_by_id: pd.Series,
                            seed_names: set[str], overrides: pd.DataFrame,
                            tol: float) -> gpd.GeoDataFrame:
    """Footprint blocks with `reason`: seed | touching_unnamed | override_include.

    blocks: merged blocks (one row per BUILDINGID). overrides: rows with buildingid and
    action ('include' | 'exclude' | 'seed') for this asset; 'seed' blocks are seeds whatever
    their name (hand footprints, decision 2026-10-07).
    """
    names = blocks.BUILDINGID.map(names_by_id)
    named = names.notna()
    is_seed = names.apply(lambda n: isinstance(n, tuple) and bool(seed_names & set(n)))
    is_seed |= blocks.BUILDINGID.isin(set(overrides.loc[overrides.action == "seed", "buildingid"]))
    if not is_seed.any():
        raise ValueError(f"no iB1000 block named {sorted(seed_names)}")
    seed_union = unary_union(blocks.geometry[is_seed].values)
    touching = ~named & (blocks.distance(seed_union) <= tol)
    reason = pd.Series(None, index=blocks.index, dtype=object)
    reason[touching] = "touching_unnamed"
    reason[is_seed] = "seed"
    include = set(overrides.loc[overrides.action == "include", "buildingid"])
    exclude = set(overrides.loc[overrides.action == "exclude", "buildingid"])
    reason[blocks.BUILDINGID.isin(include) & reason.isna()] = "override_include"
    reason[blocks.BUILDINGID.isin(exclude)] = None
    out = blocks.assign(reason=reason, names=names.apply(lambda n: " / ".join(n) if isinstance(n, tuple) else None))
    return out[out.reason.notna()].copy()


def qa_suggestion(selected: pd.DataFrame, above_seed_m: float) -> pd.Series:
    """QA suggestion for touching unnamed blocks: exclude | review | include ('' otherwise).

    'exclude' when BASELEVEL > seeds' min BASELEVEL + above_seed_m (sits on the podium).
    """
    seed_min = selected.loc[selected.reason == "seed", "BASELEVEL"].min()
    def one(r):
        if r.reason != "touching_unnamed":
            return ""
        if pd.isna(r.BASELEVEL) or pd.isna(seed_min):
            return "review"
        return "exclude" if r.BASELEVEL > seed_min + above_seed_m else "include"
    return selected.apply(one, axis=1)


def touching_reason(geom, base, roof, area: float, suggestion: str, seed_geom, seed_min,
                    no_levels_max_area_m2: float, above_seed_m: float) -> str:
    """One line on why a touching unnamed block got its QA suggestion, for the review list:
    its levels against the seeds' min BASELEVEL (or why it has none), how it touches the
    seed (shared edge length, podium under the seed, or mostly inside the seed outline)."""
    na = lambda x: x is None or pd.isna(x)
    if na(base) and na(roof):
        lv = f"no BASELEVEL/ROOFLEVEL, {area:.0f} m² (not < {no_levels_max_area_m2:g} m², so rule_v1 left it)"
    elif na(base):
        lv = f"no BASELEVEL (roof {roof:.1f} mPD)"
    elif na(seed_min):
        lv = f"base {base:.1f} mPD; seeds have no BASELEVEL to compare"
    else:
        diff = base - seed_min
        if abs(diff) <= 0.5:
            where = "at the seed's ground level"
        elif diff < 0:
            where = f"{-diff:.1f} m below the seed's lowest base (lower deck or podium)"
        else:
            where = f"{diff:.1f} m above the seed's lowest base (within +{above_seed_m:g} m)"
        lv = f"base {base:.1f} mPD, {where}"
    inside = geom.intersection(seed_geom).area / area if area else 0.0
    shared = geom.boundary.intersection(seed_geom.buffer(0.1)).length  # exact neighbours
    if shared < 2:  # separated by a small gap (touch tolerance 0.5 m)
        shared = geom.boundary.intersection(seed_geom.buffer(0.5)).length
    if inside > 0.5:
        rel = f"{inside:.0%} inside the seed outline (structure over/under the seed)"
    elif not na(roof) and not na(seed_min) and abs(roof - seed_min) <= 1.0:
        rel = f"roof {roof:.1f} mPD ≈ seed base: podium under the seed; shares {shared:.0f} m of edge"
    elif shared >= 2:
        rel = f"shares {shared:.0f} m of edge with the seed"
    else:
        rel = "touches the seed at a point (< 2 m of edge)"
    return f"{suggestion}: {lv}; {rel}"


def touches_outer_edge(geom, sheets_extent, tol: float) -> bool:
    """True if geom lies within tol of the outer boundary of the loaded sheets."""
    return bool(geom.distance(sheets_extent.boundary) <= tol)


def dissolve_sites(sites: gpd.GeoDataFrame) -> gpd.GeoDataFrame:
    """One row per SITEID (pieces clipped at sheet edges unioned)."""
    return sites.dissolve(by="SITEID", aggfunc={"SITECODE": "first"}).reset_index()


def site_name_table(site_names: pd.DataFrame, relate: pd.DataFrame) -> pd.DataFrame:
    """SITEID -> ENGLISHNAME rows (all sheets)."""
    return relate[["SITENAMEID", "SITEID"]].drop_duplicates().merge(
        site_names.drop_duplicates("SITENAMEID")[["SITENAMEID", "ENGLISHNAME"]], on="SITENAMEID")


def estate_sites(footprint, sites: gpd.GeoDataFrame, n: pd.DataFrame) -> str:
    """Names (and SITECODE) of iB1000 Sites overlapping the footprint; context only.
    sites: dissolve_sites output; n: site_name_table output."""
    s = sites.iloc[sites.sindex.query(footprint, predicate="intersects")]
    s = s[s.intersection(footprint).area > 0]
    out = []
    for r in s.itertuples():
        nm = sorted(set(n.loc[n.SITEID == r.SITEID, "ENGLISHNAME"].dropna()))
        out.append(f"{' / '.join(nm) or 'unnamed'} ({r.SITECODE})")
    return "; ".join(sorted(out))


# --- elevation ---------------------------------------------------------------------------

def ring(footprint, width_m: float):
    """Ring of width_m outside the footprint."""
    return footprint.buffer(width_m).difference(footprint)


def masked_values(arr: np.ndarray, transform, geom) -> tuple[np.ndarray, int]:
    """Values of pixels whose centre is in geom (NaN = nodata dropped), and the nodata count."""
    inside = ~geometry_mask([geom], out_shape=arr.shape, transform=transform, all_touched=False)
    v = arr[inside]
    n_nodata = int(np.isnan(v).sum())
    return v[~np.isnan(v)], n_nodata


def entry_level(values: np.ndarray, percentile: float) -> float:
    """Percentile of ring ground levels (linear interpolation); NaN if no values."""
    if values.size == 0:
        return float("nan")
    return float(np.percentile(values, percentile, method="linear"))


def podium_flag(min_baselevel: float, ring_median: float, threshold_m: float) -> bool:
    """True if the lowest block base sits more than threshold_m above the ring median.
    False when either level is missing (reported separately)."""
    if pd.isna(min_baselevel) or pd.isna(ring_median):
        return False
    return bool(min_baselevel - ring_median > threshold_m)


def slope_flag(ring_low: float, ring_high: float, threshold_m: float) -> bool:
    """True if the ring's high percentile exceeds its low percentile by more than threshold_m.
    False when either level is missing."""
    if pd.isna(ring_low) or pd.isna(ring_high):
        return False
    return bool(ring_high - ring_low > threshold_m)


def footprint_levels(sel: gpd.GeoDataFrame, arr: np.ndarray, transform, el: dict):
    """Footprint (union of the FINAL blocks, after overrides), its ring and level columns.

    Returns (footprint, ring geometry, dict of columns). el = config `elevation`.
    """
    fp = unary_union(sel.geometry.values)
    rg = ring(fp, el["ring_width_m"])
    v, n_nodata = masked_values(arr, transform, rg)
    lo_pct, hi_pct = el["slope_flag_percentiles"]
    lo, hi = entry_level(v, lo_pct), entry_level(v, hi_pct)
    median = entry_level(v, 50)
    min_base = sel.BASELEVEL.min()
    min_base = float("nan") if pd.isna(min_base) else float(min_base)
    r2 = lambda x: None if pd.isna(x) else round(float(x), 2)
    cols = {
        "entry_level_mpd": r2(entry_level(v, el["entry_level_percentile"])),
        "ring_median_mpd": r2(median),
        "ring_min_mpd": r2(v.min()) if v.size else None,
        f"ring_p{hi_pct}_mpd": r2(hi),
        "slope_flag": slope_flag(lo, hi, el["slope_flag_threshold_m"]),
        "ring_px": int(v.size), "ring_nodata_px": n_nodata,
        "ring_outside_lidar": not raster_bounds(arr, transform).contains(rg),
        "min_baselevel_mpd": r2(min_base),
        "lowest_floor_mpd": r2(min_base),
        "podium_flag": podium_flag(min_base, median, el["podium_flag_threshold_m"]),
    }
    return fp, rg, cols


def point_buffer_levels(pts, arr: np.ndarray, transform, radius_m: float, el: dict):
    """Approximate levels for an asset without a footprint (footprint_method
    'point_buffer_approx'): LiDAR pixels (centre inside) within radius_m of any accepted
    point. entry_level_mpd = entry_level_percentile (p10), ring_median_mpd = median of
    that circle (not a ring). Returns (circle geometry, dict of columns)."""
    circle = unary_union([p.buffer(radius_m) for p in pts])
    v, n_nodata = masked_values(arr, transform, circle)
    r2 = lambda x: None if pd.isna(x) else round(float(x), 2)
    return circle, {
        "entry_level_mpd": r2(entry_level(v, el["entry_level_percentile"])),
        "ring_median_mpd": r2(entry_level(v, 50)),
        "ring_px": int(v.size), "ring_nodata_px": n_nodata,
        "ring_outside_lidar": not raster_bounds(arr, transform).contains(circle),
    }


def dtm5m_levels(pts, dtm_path: Path, radius_m: float, el: dict) -> dict:
    """Fallback levels from the LandsD 5 m DTM (int32 whole metres, sea = 0, +-5 m) within
    radius_m of the point(s), for assets in elevation.dtm_5m_fallback_valuation_nos whose
    2020 LiDAR circle has no data (decision 2026-10-06; the 5 m DTM is otherwise never used
    for flood depth). Sea pixels (0) are not ground and are dropped; if only sea remains the
    levels are missing. level_source 'dtm_5m_fallback', level_confidence 'low'."""
    circle = unary_union([p.buffer(radius_m) for p in pts])
    with rasterio.open(dtm_path) as src:
        win = rasterio.windows.from_bounds(*circle.buffer(2 * abs(src.res[0])).bounds, src.transform)
        win = win.round_offsets().round_lengths()
        arr = src.read(1, window=win).astype("float64")
        trf, nodata = src.window_transform(win), src.nodata
    if nodata is not None:
        arr[arr == nodata] = np.nan
    v, n_nodata = masked_values(arr, trf, circle)
    sea = int((v == 0).sum())
    v = v[v != 0]
    r2 = lambda x: None if pd.isna(x) else round(float(x), 2)
    return {"entry_level_mpd": r2(entry_level(v, el["entry_level_percentile"])),
            "ring_median_mpd": r2(entry_level(v, 50)), "ring_px": int(v.size),
            "ring_nodata_px": n_nodata, "dtm5m_sea_px": sea,
            "level_source": "dtm_5m_fallback", "level_confidence": "low"}


def lidar_mosaic(sheet_nos: list[str], lidar_dir: Path):
    """Mosaic of the LiDAR tiles for these sheets: (array float64 with NaN, transform)."""
    files = []
    for s in sheet_nos:
        found = sorted(lidar_dir.glob(f"{s.replace('-', '')}(*/*.tif"))
        if len(found) != 1:
            raise FileNotFoundError(f"expected 1 LiDAR tile for sheet {s}, found {len(found)}")
        files += found
    srcs = [rasterio.open(f) for f in files]
    try:
        arr, transform = merge(srcs)
        nodata = srcs[0].nodata
    finally:
        for s in srcs:
            s.close()
    arr = arr[0].astype("float64")
    arr[arr == nodata] = np.nan
    return arr, transform


def lidar_tile_name(sheet_nos):
    """LiDAR DTM_2020 TILE_NAME for iB1000 sheet numbers: same grid, dashes dropped."""
    return pd.Series(sheet_nos).str.replace("-", "", regex=False).tolist()


def read_lidar_index(path: Path = LIDAR_INDEX) -> pd.DataFrame:
    """DTM_2020 grid index attributes (TILE_NAME, Filename, URL). Tile geometry comes from
    the iB1000 tile index (same grid, exact HK1980), not from this WGS84 file."""
    return pyogrio.read_dataframe(path, read_geometry=False)[["TILE_NAME", "Filename", "URL"]]


def raster_bounds(arr: np.ndarray, transform):
    h, w = arr.shape
    x0, y0 = transform * (0, 0)
    x1, y1 = transform * (w, h)
    return box(min(x0, x1), min(y0, y1), max(x0, x1), max(y0, y1))


# --- triage ------------------------------------------------------------------------------

# Reasons kept as columns but not on the review list (decision 2026-10-06)
INFO_REASONS = ("elevation:slope", "elevation:district_outlier", "geocode:low_score")
# Reasons that need a user decision (a file to edit), as opposed to a map check
DECISION_REASONS = ("footprint:no_named_seed", "footprint:touching_unnamed_undecided",
                    "footprint:rule_v2_now_below_threshold", "elevation:no_level")


def district_ranges(levels: pd.DataFrame, k: float, min_n: int) -> pd.DataFrame:
    """Per district typical entry-level range from our own assets (Tukey fences).

    levels: district, entry_level_mpd. Returns n, q1, q3, lo = q1 - k*IQR, hi = q3 + k*IQR
    (linear percentiles), for districts with at least min_n non-missing levels.
    """
    rows = []
    for d, g in levels.dropna(subset=["district", "entry_level_mpd"]).groupby("district"):
        v = g.entry_level_mpd.to_numpy(dtype=float)
        if len(v) < min_n:
            continue
        q1, q3 = np.percentile(v, [25, 75], method="linear")
        rows.append({"district": d, "n": len(v), "q1": q1, "q3": q3,
                     "lo": q1 - k * (q3 - q1), "hi": q3 + k * (q3 - q1)})
    return pd.DataFrame(rows, columns=["district", "n", "q1", "q3", "lo", "hi"]).set_index("district")


def undecided_touching(asset_blocks: pd.DataFrame, asset_overrides: pd.DataFrame) -> list[int]:
    """BUILDINGIDs of touching unnamed blocks with no include/exclude row, unless the asset
    has an 'accept' row (footprint signed off as is)."""
    if asset_overrides.action.isin(["accept", "accept_location"]).any():
        return []
    decided = set(asset_overrides.buildingid.dropna())
    t = asset_blocks[(asset_blocks.reason == "touching_unnamed") & ~asset_blocks.BUILDINGID.isin(decided)]
    return sorted(int(i) for i in t.BUILDINGID)


def footprint_triage(f: dict, undecided: list[int], suggestions: dict, ranges: pd.DataFrame,
                     als_max_m: float) -> list[tuple[str, str]]:
    """(reason, note) pairs for one footprint row (dict); note may be ''. Includes
    INFO_REASONS; asset_triage splits them off."""
    out = []
    lvl, d = f.get("entry_level_mpd"), f.get("district")
    if lvl is not None and not pd.isna(lvl) and d in ranges.index:
        r = ranges.loc[d]
        if not (r.lo <= lvl <= r.hi):
            out.append(("elevation:district_outlier",
                        f"entry {lvl} mPD outside {d} range [{r.lo:.1f}, {r.hi:.1f}] (n={int(r.n)})"))
    if f["footprint_status"] == "point_only":
        return []
    if f["footprint_status"] == "no_named_seed":
        n = f.get("n_seed_suggestions") or 0
        return [("footprint:no_named_seed",
                 f"seed names tried: {f['seed_names_tried']}; {int(n)} open seed suggestion(s) "
                 f"within {f.get('seed_search_m')} m (seed_suggestions.csv)")] + out
    rev = []
    if undecided:
        rev.append(("footprint:touching_unnamed_undecided", "undecided touching unnamed: " + ", ".join(
            f"{i} ({suggestions.get(i, '')})" for i in undecided)))
    if f["edge_touch"]:
        rev.append(("footprint:edge_touch", ""))
    if f["als_max_distance_m"] > als_max_m:
        rev.append(("geocode:als_far_from_footprint", f"ALS point {f['als_max_distance_m']:.0f} m from footprint"))
    if f["slope_flag"]:
        rev.append(("elevation:slope", f"ring p10 {f['entry_level_mpd']} / p90 {f['ring_p90_mpd']} mPD"))
    if f["podium_flag"]:
        rev.append(("elevation:podium",
                    f"min BASELEVEL {f['min_baselevel_mpd']} vs ring median {f['ring_median_mpd']} mPD"))
    if f["ring_outside_lidar"]:
        rev.append(("elevation:ring_outside_lidar", ""))
    return rev + out


def review_status(status: str, review: list[str], priority: str, n_suggestions, qa_reviewed: bool,
                  has_rule_v2: bool, accepted: bool = False, location_approx: bool = False) -> str:
    """Where an asset stands (decision 2026-10-06), in order:
    point_only               no. 130
    location_approximate     an 'accept_location' row: the point itself is approximate (no
                             block identified) and the result is insensitive to it (decision
                             2026-10-07, reason in the row's evidence); off every open list
    footprint_unconfirmed    no named seed, but an 'accept' row in footprint_overrides.csv: the
                             approximate location is kept and the footprint search closed
                             because the result is insensitive to it (decision 2026-10-07,
                             reason in the row's evidence); off every open list
    review                   on the review list: priority A with a review reason
    approx_pending_week3     priority A, no named seed and no open seed suggestion (none,
                             or all marked 'no'), with an approximate level (not
                             elevation:no_level): accepted for now; Week 3 flags any within 1 m of a
                             modelled flood level for a hand-drawn footprint
    auto_accepted_hillside   priority B (entry >= threshold) with a review reason, or with
                             touching blocks excluded by rule_v2_hillside
    hand_reviewed            no review reason left; the user has decision rows
    accepted_clean           no review reason; matched automatically
    """
    if status == "point_only":
        return "point_only"
    if location_approx:
        return "location_approximate"
    if status == "no_named_seed" and accepted:
        return "footprint_unconfirmed"
    if priority == "A" and review:
        if status == "no_named_seed" and not n_suggestions and "elevation:no_level" not in review:
            return "approx_pending_week3"
        return "review"
    if priority == "B" and (review or has_rule_v2):
        return "auto_accepted_hillside"
    return "hand_reviewed" if qa_reviewed else "accepted_clean"


def review_priority(level, below_mpd: float) -> str:
    """'A' if the entry level is below below_mpd or unknown (conservative), else 'B'."""
    return "A" if level is None or pd.isna(level) or level < below_mpd else "B"


TRIAGE_COLS = ["priority", "review_status", "valuation_no", "name", "district", "footprint_status",
               "footprint_method", "entry_level_mpd", "level_source", "value_hkd", "needs_decision",
               "reasons", "notes", "info_flags", "info_notes", "qa_reviewed", "als_source"]


def asset_triage(footprints: pd.DataFrame, blocks: pd.DataFrame, points: pd.DataFrame,
                 overrides: pd.DataFrame, ranges: pd.DataFrame, als_max_m: float,
                 values: pd.Series, priority_below_mpd: float) -> pd.DataFrame:
    """One row per asset with priority and review_status (see review_status); the review
    list is the rows with review_status 'review'. Reasons outside INFO_REASONS are review
    reasons; informational ones go to info_flags / info_notes. An asset with
    rule_v2_hillside rows that is now priority A (entry level fell below the threshold
    after the exclusions) gets footprint:rule_v2_now_below_threshold. Sorted by entry level
    ascending (approximate for unmatched assets; unknown last), then value descending."""
    rule_v2 = set(overrides.loc[overrides.decided_by.fillna("") == "rule_v2_hillside", "valuation_no"])
    geo = geocode_triage(points).set_index("valuation_no")
    rows = []
    for f in footprints.to_dict("records"):
        no = f["valuation_no"]
        b = blocks[blocks.valuation_no == no]
        sugg = dict(zip(b.BUILDINGID.astype(int), b.suggestion))
        und = undecided_touching(b, overrides[overrides.valuation_no == no])
        pairs = footprint_triage(f, und, sugg, ranges, als_max_m)
        if no in geo.index and f["footprint_status"] != "point_only":
            pairs = [(r, "") for r in geo.loc[no, "reasons"].split(";")] + pairs
        lvl = f.get("entry_level_mpd")
        prio = review_priority(lvl, priority_below_mpd) if f["footprint_status"] != "point_only" else None
        if prio is not None and (lvl is None or pd.isna(lvl)):
            pairs = [("elevation:no_level",
                      f"no ground level ({f.get('level_source')}): check the point / coordinates")] + pairs
        if no in rule_v2 and prio == "A":
            pairs = [("footprint:rule_v2_now_below_threshold",
                      f"touching blocks excluded by rule_v2_hillside, entry level now {lvl} mPD")] + pairs
        review = [(r, n) for r, n in pairs if r not in INFO_REASONS]
        info = [(r, n) for r, n in pairs if r in INFO_REASONS]
        status = review_status(f["footprint_status"], [r for r, _ in review], prio,
                               f.get("n_seed_suggestions"), bool(f["qa_reviewed"]), no in rule_v2,
                               bool((overrides.loc[overrides.valuation_no == no, "action"] == "accept").any()),
                               bool((overrides.loc[overrides.valuation_no == no, "action"] == "accept_location").any()))
        rows.append({
            "priority": prio, "review_status": status, "valuation_no": no, "name": f["name"],
            "district": f.get("district"), "footprint_status": f["footprint_status"],
            "footprint_method": f.get("footprint_method"), "entry_level_mpd": lvl,
            "level_source": f.get("level_source"),
            "value_hkd": values.get(no), "needs_decision": any(r in DECISION_REASONS for r, _ in review),
            "reasons": ";".join(r for r, _ in review), "notes": " | ".join(n for _, n in review if n),
            "info_flags": ";".join(r for r, _ in info), "info_notes": " | ".join(n for _, n in info if n),
            "qa_reviewed": f["qa_reviewed"], "als_source": f["als_source"]})
    t = pd.DataFrame(rows, columns=TRIAGE_COLS)
    t["_lvl"] = pd.to_numeric(t.entry_level_mpd, errors="coerce")
    t = t.sort_values(["_lvl", "value_hkd"], ascending=[True, False], na_position="last")
    return t.drop(columns="_lvl").reset_index(drop=True)


def asset_districts(points: pd.DataFrame) -> pd.Series:
    """valuation_no -> district of the asset's accepted point(s) (';'-joined if they differ;
    '' if none). Manual points carry only a district from their manual source (decision
    2026-10-06), never the cached ALS district of the asset's query."""
    return points.groupby("valuation_no").district.agg(
        lambda x: ";".join(sorted({d for d in x if isinstance(d, str) and d})))


TOUCHING_REVIEW_COLS = ["valuation_no", "name", "buildingid", "type", "base_mpd", "roof_mpd", "area_m2",
                        "suggestion", "reason", "decision"]


def touching_review(blocks: gpd.GeoDataFrame, overrides: pd.DataFrame, review_nos: set,
                    rules_cfg: dict) -> pd.DataFrame:
    """Undecided touching unnamed blocks of review-list assets, one row each, with the QA
    suggestion and a one-line reason (touching_reason); decision blank."""
    rows = []
    for no in sorted(review_nos):
        b = blocks[blocks.valuation_no == no]
        seeds = b[b.reason == "seed"]
        if seeds.empty:
            continue
        seed_geom = unary_union(seeds.geometry.values)
        seed_min = seeds.BASELEVEL.min()
        und = set(undecided_touching(b, overrides[overrides.valuation_no == no]))
        for r in b[b.BUILDINGID.isin(und)].itertuples():
            rows.append({
                "valuation_no": no, "name": r.name, "buildingid": int(r.BUILDINGID),
                "type": r.TYPEOFBUILDINGBLOCK, "base_mpd": r.BASELEVEL, "roof_mpd": r.ROOFLEVEL,
                "area_m2": r.area_m2, "suggestion": r.suggestion,
                "reason": touching_reason(r.geometry, r.BASELEVEL, r.ROOFLEVEL, r.area_m2, r.suggestion,
                                          seed_geom, seed_min, rules_cfg["no_levels_max_area_m2"],
                                          rules_cfg["above_seed_min_m"]),
                "decision": ""})
    return pd.DataFrame(rows, columns=TOUCHING_REVIEW_COLS)


def keep_touching_decisions(new: pd.DataFrame, prev: pd.DataFrame | None) -> pd.DataFrame:
    """Carry decisions from a previous touching_review.csv (valuation_no + buildingid)."""
    if prev is None or prev.empty:
        return new
    d = {(int(n), int(b)): str(x) for n, b, x in zip(prev.valuation_no, prev.buildingid, prev.decision)
         if str(x).strip()}
    out = new.copy()
    out["decision"] = [d.get((int(n), int(b)), "") for n, b in zip(out.valuation_no, out.buildingid)]
    return out


# --- main --------------------------------------------------------------------------------

FOOTPRINT_COLS = ["valuation_no", "name", "district", "footprint_status", "footprint_method",
                  "seed_source", "seed_names", "seed_search_m", "n_seed_suggestions", "n_blocks",
                  "block_ids", "n_touching_unnamed", "footprint_area_m2_qa", "edge_touch", "sheets",
                  "sheet_revision_dates", "lidar_tiles", "estate_sites"]
FLAG_COLS = ["edge_touch", "slope_flag", "ring_outside_lidar", "podium_flag", "qa_reviewed"]
SUGGESTION_COLS = ["valuation_no", "name", "priority", "review_status", "token", "token_source",
                   "candidate_name", "buildingids", "sheets", "distance_m", "keyword", "decision"]


def _on_disk(sheet: str) -> bool:
    return (SHEETS_DIR / sheet / f"{sheet}.gdb").exists()


def _tiles_for(geom, tile_index: gpd.GeoDataFrame) -> list[str]:
    return sorted(tile_index[tile_index.intersects(geom)].SHEETNO)


def run(pilot: bool) -> None:
    cfg = load_config("assumptions")
    crs, fp_cfg, el, tr_cfg = cfg["crs"]["analysis"], cfg["footprint"], cfg["elevation"], cfg["triage"]

    assets = read_hk_asset_names(XLSX).sort_values("no")
    if pilot:
        missing = set(cfg["pilot"]["valuation_nos"]) - set(assets.no)
        if missing:
            raise ValueError(f"pilot valuation nos. not in {XLSX.name}: {sorted(missing)}")
        assets = assets[assets.no.isin(cfg["pilot"]["valuation_nos"])]
    point_only = set(fp_cfg["point_only_valuation_nos"])
    values = read_hk_asset_values(XLSX)

    aliases = pd.read_csv(MANUAL / "asset_name_aliases.csv")
    overrides = pd.read_csv(MANUAL / "footprint_overrides.csv", dtype={"buildingid": "Int64"})
    bad = set(overrides.action) - {"include", "exclude", "accept", "accept_location", "seed"}
    if bad:
        raise ValueError(f"unknown action(s) in footprint_overrides.csv: {sorted(bad)}")
    reviewed = set(overrides.loc[~overrides.decided_by.fillna("").str.startswith("rule_"), "valuation_no"])

    points = load_accepted_points(assets, INTERIM / "geocode_accepted.csv", crs)
    shared = shared_point_nos(points)
    points["shared_point"] = points.valuation_no.isin(shared)
    districts = asset_districts(points)
    tile_index = pyogrio.read_dataframe(TILE_INDEX, layer="TileIndex")
    footprinted = assets[~assets.no.isin(point_only)]
    xy = {a.no: points.loc[points.valuation_no == a.no, ["easting", "northing"]].values
          for a in footprinted.itertuples()}
    # required: sheets within sheet_search_buffer_m; extended seed search (only for assets
    # without a seed there): sheets within seed_search_radius_m, listed if not downloaded
    sheets = {no: sheets_for_points(p, tile_index, fp_cfg["sheet_search_buffer_m"], SHEETS_DIR)
              for no, p in xy.items()}
    sheets_ext = {no: sheets_for_points(p, tile_index, fp_cfg["seed_search_radius_m"])
                  for no, p in xy.items()}
    all_sheets = pd.concat([*sheets.values(), *sheets_ext.values()]).drop_duplicates("sheet_no")
    all_sheets = all_sheets[all_sheets.sheet_no.map(_on_disk)].sort_values("sheet_no")

    extract = INTERIM / "ib1000_extract.gpkg"
    print(f"extracting {len(all_sheets)} iB1000 sheets -> {extract.relative_to(ROOT)}")
    extract_ib1000(all_sheets, SHEETS_DIR, extract, crs)
    building = gpd.GeoDataFrame(read_extract(extract, "Building"), crs=crs)
    blocks_all = merge_sheet_pieces(building, "BUILDINGID", BLOCK_ATTRS)
    loaded_extent = unary_union(tile_index[tile_index.SHEETNO.isin(all_sheets.sheet_no)].geometry.values)
    names_by_id = building_names(read_extract(extract, "BUILDINGNAME"),
                                 read_extract(extract, "BUILDINGRELATEBUILDINGNAME"))
    sites = dissolve_sites(gpd.GeoDataFrame(read_extract(extract, "Site"), crs=crs))
    site_names = site_name_table(read_extract(extract, "SITENAME"), read_extract(extract, "SITERELATESITENAME"))
    sheet_geom = tile_index.set_index("SHEETNO").geometry

    def blocks_in(s_nos):
        area = unary_union(sheet_geom.loc[s_nos].values)
        return blocks_all.iloc[blocks_all.sindex.query(area, predicate="intersects")]

    # Pass 1: footprints (no LiDAR needed)
    fp_rows, sel_by_no, sugg_rows, missing_ext = [], {}, [], {}
    for a in assets.itertuples():
        pts = points[points.valuation_no == a.no]
        base = {"valuation_no": a.no, "name": a.name, "district": districts.get(a.no),
                "qa_reviewed": a.no in reviewed,
                "als_source": ",".join(sorted(set(pts.als_source))),
                "geocode_flags": ";".join(sorted({f for s in pts["flags"].fillna("")
                                                   for f in str(s).split(";") if f}))}
        if a.no in point_only:
            fp_rows.append({**base, "footprint_status": "point_only", "geometry": None})
            continue
        tried, how = seed_candidates(a.no, a.name, aliases)
        seed_names = seed_names_for(tried, how, names_by_id)
        asset_ov = overrides[overrides.valuation_no == a.no]
        sel = None
        for radius, s_df in ((fp_cfg["sheet_search_buffer_m"], sheets[a.no]),
                             (fp_cfg["seed_search_radius_m"], sheets_ext[a.no])):
            s_nos = list(s_df.sheet_no)
            lacking = [s for s in s_nos if not _on_disk(s)]
            if lacking:
                missing_ext[a.no] = lacking
                break
            cand = blocks_in(s_nos)
            try:
                sel = select_footprint_blocks(cand, names_by_id, seed_names, asset_ov,
                                              fp_cfg["touch_tolerance_m"])
                break
            except ValueError as e:
                if not str(e).startswith("no iB1000 block named"):
                    raise
        base.update(seed_source="override_seed" if (asset_ov.action == "seed").any() else how,
                    seed_names_tried=" / ".join(tried), seed_search_m=radius,
                    sheets=",".join(s_nos), sheet_revision_dates=",".join(s_df.sheet_revision_date))
        if sel is None:
            tok, tok_src = asset_token(a.no, a.name, aliases)
            sg = seed_suggestions(cand if not lacking else blocks_in([s for s in s_nos if _on_disk(s)]),
                                  names_by_id, list(pts.geometry), tok, fp_cfg["seed_search_radius_m"])
            sugg_rows.append(sg.assign(valuation_no=a.no, name=a.name, token=tok, token_source=tok_src))
            circle = unary_union([p.buffer(el["approx_point_buffer_m"]) for p in pts.geometry])
            fp_rows.append({**base, "footprint_status": "no_named_seed",
                            "footprint_method": "point_buffer_approx", "n_seed_suggestions": len(sg),
                            "lidar_tiles": ",".join(_tiles_for(circle, tile_index)), "geometry": None})
            continue
        sel["suggestion"] = qa_suggestion(sel, fp_cfg["suggest_exclude_above_seed_m"])
        sel["edge_touch"] = [touches_outer_edge(g, loaded_extent, fp_cfg["edge_tolerance_m"])
                             for g in sel.geometry]
        fp = unary_union(sel.geometry.values)
        sel_by_no[a.no] = sel
        fp_rows.append({
            **base, "footprint_status": "matched", "footprint_method": "ib1000_footprint",
            "seed_names": " / ".join(sorted({n.strip() for t in sel.loc[sel.reason == "seed", "names"]
                                             if isinstance(t, str)  # 'seed' override blocks may be unnamed
                                             for n in t.split(" / ") if n in seed_names})),
            "n_blocks": len(sel), "block_ids": ",".join(map(str, sorted(sel.BUILDINGID))),
            "n_touching_unnamed": int((sel.reason == "touching_unnamed").sum()),
            "footprint_area_m2_qa": round(fp.area, 1), "edge_touch": bool(sel.edge_touch.any()),
            "lidar_tiles": ",".join(_tiles_for(ring(fp, el["ring_width_m"]).union(fp), tile_index)),
            "estate_sites": estate_sites(fp, sites, site_names),
            "als_max_distance_m": round(float(pts.distance(fp).max()), 1),
            "geometry": fp,
        })

    # iB1000 sheets for the extended seed search: list them; stop if any is not downloaded
    ext = sorted({s for v in missing_ext.values() for s in v})
    pd.DataFrame({"sheet_no": ext, "for_valuation_nos": [
        ",".join(str(n) for n, v in missing_ext.items() if s in v) for s in ext]}).to_csv(
        INTERIM / "ib1000_sheets_needed.csv", index=False)
    if ext:
        raise FileNotFoundError(
            f"{len(ext)} iB1000 sheets for the {fp_cfg['seed_search_radius_m']} m seed search not "
            f"downloaded ({len(missing_ext)} assets); listed in data/interim/ib1000_sheets_needed.csv. "
            "Run scripts/fetch_terrain.py --ib1000 --delete-zips, then rerun.")

    # LiDAR tiles under footprint + ring (or the approx circle): list; stop if any is missing
    need = sorted({t for r in fp_rows for t in (r.get("lidar_tiles") or "").split(",") if t})
    have = [len(list(LIDAR_DIR.glob(f"{t}(*/*.tif"))) == 1 for t in lidar_tile_name(need)]
    pd.DataFrame({"sheet_no": need, "tile": lidar_tile_name(need), "have": have}).to_csv(
        INTERIM / "lidar_tiles_needed.csv", index=False)
    if not all(have):
        raise FileNotFoundError(
            f"{have.count(False)} of {len(need)} LiDAR tiles not downloaded; listed in "
            "data/interim/lidar_tiles_needed.csv. Run scripts/fetch_terrain.py --lidar, then rerun.")

    # Pass 2: levels from the final footprint, or approximate levels around the point
    block_rows = []
    for r in fp_rows:
        if r["footprint_status"] == "point_only":
            continue
        arr, trf = lidar_mosaic(r["lidar_tiles"].split(","), LIDAR_DIR)
        if r["footprint_status"] == "no_named_seed":
            pts = list(points[points.valuation_no == r["valuation_no"]].geometry)
            _, levels = point_buffer_levels(pts, arr, trf, el["approx_point_buffer_m"], el)
            r.update(levels, level_source="lidar_2020", level_confidence="normal")
            if r["entry_level_mpd"] is None and r["valuation_no"] in el["dtm_5m_fallback_valuation_nos"]:
                r.update(dtm5m_levels(pts, DTM_5M, el["approx_point_buffer_m"], el))
            continue
        sel = sel_by_no[r["valuation_no"]]
        _, _, levels = footprint_levels(sel, arr, trf, el)
        r.update(levels, level_source="lidar_2020", level_confidence="normal")
        b = sel.drop(columns=["pieces"]).assign(valuation_no=r["valuation_no"], name=r["name"])
        b["height_m"] = (b.ROOFLEVEL - b.BASELEVEL).round(2)
        b["area_m2"] = b.area.round(1)
        block_rows.append(b)

    footprints = gpd.GeoDataFrame(pd.DataFrame(fp_rows), geometry="geometry", crs=crs)
    level_cols = [c for c in footprints.columns if c not in FOOTPRINT_COLS + ["geometry"]]
    footprints = footprints.reindex(columns=FOOTPRINT_COLS + level_cols + ["geometry"])
    footprints.insert(footprints.columns.get_loc("geometry"), "shared_point",
                      footprints.valuation_no.isin(shared))
    for c in FLAG_COLS:  # bools with nulls (unmatched assets) must not be written as strings
        footprints[c] = footprints[c].astype("boolean")
    blocks_out = gpd.GeoDataFrame(pd.concat(block_rows, ignore_index=True), crs=crs)
    blocks_out["LASTUPDATEDATE"] = blocks_out.LASTUPDATEDATE.astype(str)
    blocks_out = blocks_out[["valuation_no", "name", "BUILDINGID", "names", "reason", "suggestion",
                             "TYPEOFBUILDINGBLOCK", "BASELEVEL", "ROOFLEVEL", "height_m", "area_m2",
                             "BASELEVELDATASOURCE", "BUILDINGSTATUS", "LASTUPDATEDATE", "sheets",
                             "edge_touch", "geometry"]].sort_values(["valuation_no", "BUILDINGID"])

    # Suggestions with the user's decisions carried over; open = not marked 'no'. The
    # status counts open suggestions only: an asset whose suggestions were all rejected
    # counts as having none (decision 2026-10-06).
    sugg = pd.concat(sugg_rows, ignore_index=True) if sugg_rows else pd.DataFrame(columns=SUGGESTION_COLS)
    sugg_path = INTERIM / "seed_suggestions.csv"
    decisions = read_seed_decisions()
    view = pd.read_csv(sugg_path, dtype=str, keep_default_na=False) if sugg_path.exists() else None
    if view is not None:
        view["valuation_no"] = view.valuation_no.astype(int)
    stray = unrecorded_view_decisions(view, decisions)
    if len(stray):
        raise ValueError(
            f"{len(stray)} decision(s) in {sugg_path.relative_to(ROOT)} are not in "
            f"{SEED_DECISIONS.relative_to(ROOT)} (the view is regenerated; decisions live in the "
            "tracked file). Add them there, then rerun: "
            + "; ".join(f"no. {r.valuation_no} {r.candidate_name!r}={r.decision}" for r in stray.itertuples()))
    sugg = merge_decisions(sugg, decisions)
    n_open = open_suggestions(sugg)
    unmatched = footprints.footprint_status == "no_named_seed"
    footprints.loc[unmatched, "n_seed_suggestions"] = footprints.loc[unmatched, "valuation_no"].map(n_open).fillna(0)
    ranges = district_ranges(footprints[["district", "entry_level_mpd"]],
                             tr_cfg["district_outlier_iqr_k"], tr_cfg["district_outlier_min_assets"])
    status = asset_triage(footprints, blocks_out, points, overrides, ranges, tr_cfg["als_max_distance_m"],
                          values, tr_cfg["priority_a_below_mpd"])
    st = status.set_index("valuation_no")
    footprints["priority"] = footprints.valuation_no.map(st.priority)
    footprints["review_status"] = footprints.valuation_no.map(st.review_status)

    for gdf, fname in [(points, "assets.geojson"), (footprints, "asset_footprints.geojson"),
                       (blocks_out, "asset_blocks.geojson")]:
        out = INTERIM / fname
        out.unlink(missing_ok=True)
        gdf.to_file(out, driver="GeoJSON")
        print(f"wrote {out.relative_to(ROOT)} ({len(gdf)} rows)")

    sugg["priority"] = sugg.valuation_no.map(st.priority)
    sugg["review_status"] = sugg.valuation_no.map(st.review_status)
    sugg = sugg[SUGGESTION_COLS].sort_values(["review_status", "valuation_no", "distance_m"],
                                             key=lambda c: c.ne("review") if c.name == "review_status" else c)
    sugg.to_csv(sugg_path, index=False)
    print(f"wrote {sugg_path.relative_to(ROOT)} ({len(sugg)} rows, {sugg.valuation_no.nunique()} assets; "
          f"on the review list: {int((sugg.review_status == 'review').sum())} rows; "
          f"decisions kept: {int((sugg.decision != '').sum())})")

    review = status[status.review_status == "review"].reset_index(drop=True)
    review.to_csv(INTERIM / "asset_triage.csv", index=False)
    tr = touching_review(blocks_out, overrides, set(review.valuation_no), cfg["footprint_rules_v1"])
    tr_path = INTERIM / "touching_review.csv"
    prev = pd.read_csv(tr_path, dtype={"decision": str}, keep_default_na=False) if tr_path.exists() else None
    tr = keep_touching_decisions(tr, prev)
    tr.to_csv(tr_path, index=False)
    print(f"wrote {tr_path.relative_to(ROOT)} ({len(tr)} blocks, {tr.valuation_no.nunique()} assets; "
          f"suggestions {tr.suggestion.value_counts().to_dict()}; decisions kept: {int((tr.decision != '').sum())})")
    status.to_csv(INTERIM / "asset_review_status.csv", index=False)
    print(f"wrote {(INTERIM / 'asset_triage.csv').relative_to(ROOT)} (review list, {len(review)} rows) "
          f"and asset_review_status.csv ({len(status)} rows)")
    print(f"wrote {extract.relative_to(ROOT)} ({len(all_sheets)} sheets)")
    summarise(footprints, status, sugg)


def summarise(footprints: pd.DataFrame, status: pd.DataFrame, sugg: pd.DataFrame) -> None:
    print("\nfootprint_status:", footprints.footprint_status.value_counts().to_dict())
    m = footprints[footprints.footprint_status == "matched"]
    print("seed_source (matched):", m.seed_source.value_counts().to_dict(),
          "| seed_search_m:", m.seed_search_m.value_counts().to_dict())
    print("level_source:", footprints.level_source.value_counts(dropna=False).to_dict())
    print("review_status:", status.review_status.value_counts().to_dict())
    review = status[status.review_status == "review"]
    counts = pd.Series([r for s in review.reasons for r in s.split(";")], dtype=str).value_counts()
    print(f"review list: {len(review)} (A {int((review.priority == 'A').sum())}; "
          f"needs a decision {int(review.needs_decision.sum())})")
    print(counts.to_string())
    s = sugg[sugg.review_status == "review"]
    print(f"seed-suggestion decisions on the review list: {len(s)} rows for {s.valuation_no.nunique()} assets")
    t = review[review.reasons.str.contains("touching_unnamed_undecided")]
    print(f"matched priority-A assets with undecided touching blocks: {len(t)}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--pilot", action="store_true", help="run the 5 pilot assets only")
    args = parser.parse_args()
    run(pilot=args.pilot)


if __name__ == "__main__":
    main()
